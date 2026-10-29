"""The input workbook: one .xlsx with Project, Measures and IGBC sheets. Problems come back as readable messages."""
from __future__ import annotations

import io
from pathlib import Path
from typing import IO

import pandas as pd
from pydantic import ValidationError

from .schemas import EXPORT_ESCALATION, PV_AREA_PER_KWP_M2, Credit, Dataset, Level, Measure, Project

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
DEMO_INPUT = DATA_DIR / "demo_input.xlsx"

# field, label, unit, meaning. Fields with unit "%" are typed as percent in the sheet and stored as fractions.
PROJECT_FIELDS = [
    ("name", "Project name", "", "Shown on every page and in the report"),
    ("city", "City", "", "Shown in the report header"),
    ("roof_area_m2", "Usable roof area", "m²", f"Caps rooftop solar at roof ÷ {PV_AREA_PER_KWP_M2:g} m² per kWp"),
    ("baseline_energy_kwh", "Baseline electricity use", "kWh/yr", "Code-minimum building; % savings for IGBC are measured against it"),
    ("baseline_water_kl", "Baseline water use", "kL/yr", "Code-minimum building; % water saving for IGBC"),
    ("study_period", "Study period (planning horizon)", "years", "Common analysis period N for every alternative (Ch 5)"),
    ("discount_rate", "MARR (discount rate, nominal)", "%", "Minimum attractive rate of return used to discount every cash flow (Ch 7)"),
    ("inflation", "General inflation", "%", "Escalates O&M and re-purchase prices; real MARR = (1+MARR)/(1+inflation) − 1 (Ch 2)"),
    ("grid_tariff", "Grid tariff, year 1", "₹/kWh", "All-in price of a kWh bought from the grid"),
    ("tariff_escalation", "Tariff escalation", "%/yr", "Expected yearly growth of the grid tariff (geometric gradient)"),
    ("tariff_escalation_sd", "Tariff escalation uncertainty (1 sd)", "%/yr", "Spread of tariff growth used by the Monte Carlo"),
    ("export_tariff", "Solar export credit", "₹/kWh", "Paid for solar kWh exported to the grid (held flat)"),
    ("water_tariff", "Water + sewage tariff, year 1", "₹/kL", "Price of a kL of water saved"),
    ("water_escalation", "Water tariff escalation", "%/yr", "Expected yearly growth of the water tariff"),
    ("pv_yield_kwh_per_kwp", "Solar specific yield", "kWh/kWp/yr", "Year-1 output of one kWp at this site"),
    ("pv_self_consumption", "Solar self-consumption share", "%", "Share of solar used on site (worth the grid tariff); the rest is exported"),
    ("grid_emission_factor", "Grid emission factor", "tCO₂/MWh", "CO₂ avoided per MWh not drawn from the grid"),
    ("target_level", "Target IGBC level", "", "Must match a level in the IGBC_levels sheet"),
    ("base_points", "IGBC points already earned by the base design", "points", "Credits the design earns without any measure below"),
]
PCT = {f for f, _, u, _ in PROJECT_FIELDS if u.startswith("%")}

MEASURE_COLUMNS = [
    ("measure_id", "ID", "", "Short unique code, e.g. PV_100"),
    ("name", "Measure", "", "Plain name"),
    ("exclusive_group", "Group", "", "Options with the same group are alternatives: at most one is chosen. Blank = independent"),
    ("credit_id", "IGBC credit", "", "Credit the measure contributes to (from IGBC_credits)"),
    ("points_fixed", "Fixed points", "points", "Only for 'fixed' credits; threshold credits are scored from the savings"),
    ("capex", "Capex (incremental)", "₹", "Price above the code-minimum option it replaces, year 0"),
    ("life_years", "Life", "years", "Service life; re-bought when it ends inside the study period, salvage counted at the end"),
    ("om_annual", "O&M, year 1 (incremental)", "₹/yr", "Extra maintenance vs baseline; escalates with inflation"),
    ("energy_saving_kwh", "Energy saved", "kWh/yr", "Electricity saved vs baseline (efficiency measures)"),
    ("kwp", "Solar size", "kWp", "Rooftop solar capacity (solar measures only)"),
    ("water_saving_kl", "Water saved", "kL/yr", "Water saved vs baseline"),
    ("degradation", "Degradation", "%/yr", "Yearly loss of savings or solar output"),
    ("part_name", "Replaceable part", "", "One major part with a shorter life (inverter, compressor…), optional"),
    ("part_cost", "Part cost", "₹", "Today's price of that part"),
    ("part_life", "Part life", "years", "How often the part is replaced"),
    ("description", "Notes / basis", "", "How the numbers were sized; source"),
]
MEASURE_PCT = {"degradation"}

CREDIT_COLUMNS = ["credit_id", "category", "name", "max_points", "rule", "basis", "thresholds", "mandatory"]
LEVEL_COLUMNS = ["level", "min_points", "rank"]

README = [
    ("What this file is", "All inputs of GreenLCC India for one building. Edit the values, save, then open the file from the tool's Home page."),
    ("Project sheet", "One row per input. Type values in the 'value' column only; % fields are typed as percent (8 = 8%)."),
    ("Measures sheet", "One row per green option. Costs and savings are INCREMENTAL over the code-minimum option it replaces "
                       "(e.g. VRF capex = VRF price − price of the split ACs it replaces)."),
    ("Groups", "Rows sharing a Group are alternatives (one HVAC system, one solar size): the optimizer picks at most one per group."),
    ("Life and part", "A measure is re-bought when its life ends inside the study period; its one replaceable part is re-bought every part life. "
                      "Unused life at the end is credited back as salvage (planning-horizon method)."),
    ("IGBC sheets", "Credits: 'fixed' credits add the Fixed points of the measures that claim them; 'threshold_pct' credits score the combined % saving "
                    "(energy_eff, energy_re or water) against thresholds written as pct:points|pct:points. Levels: points needed for each rating."),
    ("Fixed constants", f"Roof area per kWp = {PV_AREA_PER_KWP_M2:g} m²; solar export credit escalation = {EXPORT_ESCALATION:.0%}."),
    ("Status of the demo data", "Prices and tariffs are realistic estimates with sources in the Project/Measures notes; IGBC point tables are illustrative."),
]


def _fmt(err: ValidationError) -> str:
    return "; ".join(f"{'.'.join(map(str, e['loc'])) or 'row'}: {e['msg']}" for e in err.errors())


def rows_to_models(df: pd.DataFrame, model, label: str) -> tuple[list, list[str]]:
    """Validate every row of a DataFrame; return (good models, error strings)."""
    out, errors = [], []
    df = df.astype(object).where(pd.notna(df), None)
    for i, rec in enumerate(df.to_dict("records"), start=1):
        rec = {k: v for k, v in rec.items() if v is not None or k == "exclusive_group"}
        try:
            out.append(model(**rec))
        except ValidationError as e:
            errors.append(f"{label} row {i}: {_fmt(e)}")
    return out, errors


def cross_check(ds: Dataset) -> list[str]:
    errors: list[str] = []
    ids = [m.measure_id for m in ds.measures]
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        errors.append(f"Duplicate measure ID: {sorted(dup)}")
    credit_ids = {c.credit_id for c in ds.credits}
    for m in ds.measures:
        if m.credit_id and m.credit_id not in credit_ids:
            errors.append(f"Measure {m.measure_id}: unknown IGBC credit '{m.credit_id}'")
    if ds.project.target_level not in {lv.level for lv in ds.levels}:
        errors.append(f"Target level '{ds.project.target_level}' is not in the IGBC_levels sheet")
    return errors


def measures_to_df(ms: list[Measure]) -> pd.DataFrame:
    cols = [f for f, *_ in MEASURE_COLUMNS]
    return pd.DataFrame([m.model_dump() for m in ms], columns=cols)


# ---------- read ----------


def _measure_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Accept either readable headers or field names; % columns back to fractions."""
    by_label = {lab: f for f, lab, *_ in MEASURE_COLUMNS}
    df = df.rename(columns=lambda c: by_label.get(str(c).strip(), str(c).strip()))
    df = df[[c for c in df.columns if c in {f for f, *_ in MEASURE_COLUMNS}]].dropna(how="all")
    for c in MEASURE_PCT & set(df.columns):
        df[c] = pd.to_numeric(df[c], errors="coerce") / 100
    return df


def read_workbook(src: str | Path | bytes | IO[bytes]) -> tuple[Dataset | None, list[str]]:
    """Read an input workbook (or a results workbook, which carries the same sheets)."""
    try:
        xls = pd.ExcelFile(io.BytesIO(src) if isinstance(src, bytes) else src)
    except Exception as e:
        return None, [f"Not a readable Excel file: {e}"]
    need = {"Project", "Measures", "IGBC_credits", "IGBC_levels"}
    if missing := need - set(xls.sheet_names):
        return None, [f"Missing sheet(s): {', '.join(sorted(missing))}. Start from the input workbook downloaded from Home."]
    errors: list[str] = []

    pr = xls.parse("Project", dtype=object)
    raw = {}
    for _, r in pr.iterrows():
        f, v = str(r.get("field", "")).strip(), r.get("value")
        if not f or f == "nan" or v is None or (isinstance(v, float) and v != v):
            continue
        raw[f] = (float(v) / 100 if f in PCT else v)
    try:
        project = Project(**raw)
    except ValidationError as e:
        return None, [f"Project sheet: {_fmt(e)}"]

    measures, errs = rows_to_models(_measure_frame(xls.parse("Measures")), Measure, "Measures")
    errors += errs
    credits, errs = rows_to_models(xls.parse("IGBC_credits")[CREDIT_COLUMNS].dropna(how="all"), Credit, "IGBC_credits")
    errors += errs
    levels, errs = rows_to_models(xls.parse("IGBC_levels")[LEVEL_COLUMNS].dropna(how="all"), Level, "IGBC_levels")
    errors += errs
    ds = Dataset(project=project, measures=measures, credits=credits, levels=levels)
    errors += cross_check(ds)  # check the rows that did parse, so every problem shows in one pass
    return (None, errors) if errors else (ds, [])


def load_dataset(path: str | Path = DEMO_INPUT) -> tuple[Dataset | None, list[str]]:
    p = Path(path)
    if not p.exists():
        return None, [f"{p.name}: file not found (run python scripts/make_demo_input.py)"]
    return read_workbook(p)


# ---------- write ----------


def input_sheets(ds: Dataset, notes: dict[str, str] | None = None) -> dict[str, pd.DataFrame]:
    """The four input sheets plus README, as DataFrames (also reused by the results workbook)."""
    notes = notes or {}
    p = ds.project.model_dump()
    proj = pd.DataFrame([{"field": f, "input": lab, "value": (p[f] * 100 if f in PCT else p[f]), "unit": u,
                          "meaning": mean, "source / status": notes.get(f, "")} for f, lab, u, mean in PROJECT_FIELDS])
    ms = measures_to_df(ds.measures)
    for c in MEASURE_PCT:
        ms[c] = ms[c] * 100
    ms.columns = [lab for _, lab, *_ in MEASURE_COLUMNS]
    cr = pd.DataFrame([{**c.model_dump(), "thresholds": "|".join(f"{a:g}:{b:g}" for a, b in c.thresholds)} for c in ds.credits],
                      columns=CREDIT_COLUMNS)
    lv = pd.DataFrame([lv.model_dump() for lv in ds.levels], columns=LEVEL_COLUMNS)
    cols = pd.DataFrame([{"column": lab, "unit": u, "meaning": mean} for _, lab, u, mean in MEASURE_COLUMNS])
    readme = pd.concat([pd.DataFrame(README, columns=["topic", "how it works"]),
                        cols.rename(columns={"column": "topic", "meaning": "how it works"})[["topic", "how it works"]]
                        .assign(topic=lambda d: "Measures column: " + d["topic"] + (" (" + cols["unit"] + ")").where(cols["unit"] != "", ""))],
                       ignore_index=True)
    return {"README": readme, "Project": proj, "Measures": ms, "IGBC_credits": cr, "IGBC_levels": lv}


def style_sheets(wb, wide: dict[str, dict[str, int]] | None = None) -> None:
    """Green header, frozen top row, sensible widths, wrapped long text."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head = PatternFill("solid", fgColor="1E7B4F")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for c in ws[1]:
            c.fill, c.font = head, Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)
            for c in col[1:]:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.number_format = "#,##0.00" if abs(c.value) < 100 and c.value != int(c.value) else "#,##0"
                if isinstance(c.value, str) and len(c.value) > 60:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        for col, w in (wide or {}).get(ws.title, {}).items():
            ws.column_dimensions[col].width = w


INPUT_WIDTHS = {"README": {"A": 40, "B": 110}, "Project": {"A": 22, "B": 40, "E": 60, "F": 60}}


def write_workbook(ds: Dataset, notes: dict[str, str] | None = None) -> bytes:
    """The input workbook for this dataset; read_workbook(write_workbook(ds)) == ds."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in input_sheets(ds, notes).items():
            df.to_excel(xw, sheet_name=name, index=False)
        style_sheets(xw.book, INPUT_WIDTHS)
    return buf.getvalue()
